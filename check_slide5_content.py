from pptx import Presentation


def check_slide5_content():
    """Check what content was generated for slide 5"""
    print("=== CHECKING SLIDE 5 CONTENT ===\n")

    try:
        prs = Presentation("ekona_charts_demo.pptx")

        if len(prs.slides) >= 5:
            slide = prs.slides[4]  # Slide 5 (0-indexed)
            print(f"Slide 5: {slide.slide_layout.name}")

            for shape_idx, shape in enumerate(slide.shapes):
                print(f"\nShape {shape_idx}: {shape.name}")
                print(f"  Shape type: {shape.shape_type}")

                # Check if it's a chart placeholder
                if hasattr(shape, "placeholder_format"):
                    print(f"  Placeholder type: {shape.placeholder_format.type}")

                # Check text content
                if hasattr(shape, "text_frame") and shape.text_frame:
                    text = shape.text_frame.text.strip()
                    if text:
                        print(f"  Text content: {text}")
                    else:
                        print("  Text content: [EMPTY]")
        else:
            print("Slide 5 not found")

    except Exception as e:
        print(f"Error: {e}")


if __name__ == "__main__":
    check_slide5_content()
