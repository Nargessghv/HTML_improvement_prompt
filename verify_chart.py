from pptx import Presentation


def verify_chart_creation():
    """Verify that chart was actually created"""
    print("=== VERIFYING CHART CREATION ===\n")

    try:
        prs = Presentation("test_single_chart.pptx")
        print(f"Total slides: {len(prs.slides)}")

        for slide_idx, slide in enumerate(prs.slides):
            print(f"\nSlide {slide_idx + 1}: {slide.slide_layout.name}")

            for shape_idx, shape in enumerate(slide.shapes):
                print(f"  Shape {shape_idx}: {shape.name}")
                print(f"    Shape type: {shape.shape_type}")

                # Check for chart
                if hasattr(shape, "chart"):
                    print("    ✅ CHART FOUND!")
                    try:
                        chart = shape.chart
                        print(f"    Chart type: {chart.chart_type}")
                        if hasattr(chart, "chart_title") and chart.chart_title:
                            title_text = chart.chart_title.text_frame.text
                            print(f"    Chart title: {title_text}")
                        print(f"    Series count: {len(chart.plots[0].series)}")

                        # Check series colors (ekona branding)
                        for i, series in enumerate(chart.plots[0].series):
                            try:
                                if hasattr(series, "format") and hasattr(
                                    series.format, "fill"
                                ):
                                    color = series.format.fill.fore_color.rgb
                                    print(
                                        f"    Series {i+1} color: RGB({color.red}, {color.green}, {color.blue})"
                                    )
                            except:
                                print(f"    Series {i+1} color: Could not read")

                    except Exception as e:
                        print(f"    Error reading chart details: {e}")

                elif hasattr(shape, "text_frame") and shape.text_frame:
                    text = shape.text_frame.text.strip()
                    if text:
                        print(f"    Text: {text[:50]}{'...' if len(text) > 50 else ''}")

    except Exception as e:
        print(f"Error: {e}")


if __name__ == "__main__":
    verify_chart_creation()
