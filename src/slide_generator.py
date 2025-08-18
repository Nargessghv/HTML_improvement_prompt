"""
Slide Generator Module

This module creates actual PowerPoint slides using the generated content
and the template structure.
"""

import os
from typing import Any, Dict, List, Optional

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

from .chart_generator import ChartGenerator
from .content_generator import ContentGenerator
from .html_renderer import HTMLRenderer
from .icon_manager import IconManager
from .icon_selector import IconSelector
from .llm_client import SlideContent
from .markdown_formatter import MarkdownFormatter
from .layout_formatting_extractor import LayoutFormattingExtractor


class SlideGenerator:
    """Creates PowerPoint slides from generated content"""

    def __init__(self, template_path: str, openai_api_key: Optional[str] = None):
        """
        Initialize the slide generator

        Args:
            template_path: Path to the PowerPoint template file
            openai_api_key: OpenAI API key (optional, can use environment variable)
        """
        self.template_path = template_path
        self.content_generator = ContentGenerator(template_path)
        self.chart_generator = ChartGenerator()
        self.markdown_formatter = MarkdownFormatter()  # Initialize markdown formatter
        self.layout_formatter = LayoutFormattingExtractor()  # Initialize layout formatter

        # Initialize HTML renderer for custom visualizations
        try:
            self.html_renderer = HTMLRenderer()
            print("✅ HTML renderer initialized successfully")
        except RuntimeError as e:
            print(f"⚠️ HTML renderer not available: {e}")
            self.html_renderer = None

        # Initialize icon management components
        self.icon_manager = IconManager()
        self.icon_selector = IconSelector(self.icon_manager)

        # Store current topic for icon-aware population
        self._current_topic = "Presentation Topic"

    def create_presentation(
        self, topic: str, output_path: str, layout_indices: Optional[List[int]] = None
    ) -> str:
        """
        Create a complete PowerPoint presentation

        Args:
            topic: The presentation topic
            output_path: Path where to save the generated presentation
            layout_indices: Specific layout indices to use (optional)

        Returns:
            Path to the created presentation file
        """
        print(f"Creating presentation for topic: '{topic}'")

        # Generate content for all slides using the proven approach
        # The mapping fix handles custom placeholder names correctly
        slide_contents = self.content_generator.generate_slide_contents(
            topic, layout_indices
        )

        if not slide_contents:
            raise ValueError("No slide contents generated")

        # Create the presentation
        presentation = self._create_powerpoint_presentation(slide_contents)

        # Save the presentation
        full_output_path = self._ensure_output_path(output_path)
        presentation.save(full_output_path)

        print(f"Presentation saved to: {full_output_path}")

        # Print summary
        self.content_generator.print_generation_summary(topic, slide_contents)

        return full_output_path

    def create_slide_with_dynamic_content(
        self,
        presentation: Any,
        layout_index: int,
        topic: str,
        slide_spec: Optional[Any] = None,
        slide_number: int = 1,
        total_slides: int = 1,
    ) -> None:
        """
        Create a slide and generate content based on its actual placeholder names

        Args:
            presentation: PowerPoint presentation object
            layout_index: Index of the layout to use
            topic: The presentation topic
            slide_spec: Optional slide specification with title/purpose
            slide_number: Current slide number
            total_slides: Total number of slides
        """
        # Get the layout and create the slide
        layout = presentation.slide_layouts[layout_index]
        slide = presentation.slides.add_slide(layout)

        # Get the actual placeholder names from the created slide
        actual_placeholders = self._get_actual_placeholder_info(slide)

        if not actual_placeholders:
            print(
                f"  → Slide {slide_number} has no placeholders, "
                "skipping content generation"
            )
            return

        # Extract just the placeholder names for content generation
        placeholder_names = list(actual_placeholders.keys())
        print(f"  → Found actual placeholders: {placeholder_names}")

        # Generate content based on actual placeholder names
        if hasattr(self, "content_generator") and self.content_generator:
            content = self._generate_content_for_actual_placeholders(
                placeholder_names, topic, slide_spec, slide_number, total_slides
            )

            if content:
                # Populate placeholders with generated content
                self._populate_slide_with_icons(
                    slide, content, actual_placeholders, topic
                )
            else:
                print(f"  → Warning: No content generated for slide {slide_number}")
        else:
            print(
                f"  → Warning: No content generator available for "
                f"slide {slide_number}"
            )

    def _generate_content_for_actual_placeholders(
        self,
        placeholder_names: List[str],
        topic: str,
        slide_spec: Optional[Any],
        slide_number: int,
        total_slides: int,
    ) -> Optional[Dict[str, str]]:
        """
        Generate content specifically for the actual placeholder names found on a slide

        Args:
            placeholder_names: List of actual placeholder names from the slide
            topic: The presentation topic
            slide_spec: Optional slide specification
            slide_number: Current slide number
            total_slides: Total number of slides

        Returns:
            Dictionary mapping placeholder names to content
        """
        try:
            # Import here to avoid circular imports
            from .llm_client import LLMClient

            # Create a simple layout info structure for the LLM
            layout_info = {
                "name": f"Dynamic Layout {slide_number}",
                "placeholders": [{"name": name} for name in placeholder_names],
            }

            # Use LLM client to generate content
            llm_client = LLMClient()

            if slide_spec:
                slide_content = llm_client.generate_contextual_slide_content(
                    layout_info, topic, slide_spec, slide_number, total_slides
                )
            else:
                slide_content = llm_client.generate_slide_content(
                    layout_info, topic, slide_number, total_slides
                )

            return slide_content.content if slide_content else None

        except Exception as e:
            print(f"  → Error generating content: {e}")
            return None

    def _populate_slide_with_icons(
        self,
        slide,
        content: Dict[str, str],
        actual_placeholders: Dict[str, Dict[str, Any]],
        topic: str,
    ) -> None:
        """
        Populate slide with content and automatically select/insert icons

        Args:
            slide: PowerPoint slide object
            content: Dictionary mapping placeholder names to content
            actual_placeholders: Dictionary of actual placeholder info
            topic: Presentation topic for icon selection context
        """
        # Check if this slide has ICON placeholders (specifically named with "icon")
        # NOT just any picture placeholder
        icon_placeholders = {
            name: info
            for name, info in actual_placeholders.items()
            if (info["type"] == PP_PLACEHOLDER.PICTURE and "icon" in name.lower())
        }

        if icon_placeholders:
            print(f"  → Found {len(icon_placeholders)} icon placeholders")

            # Select appropriate icons based on content
            icon_selections = self.icon_selector.select_icons_for_content(
                content, topic
            )

            # Prepare the selected icons
            prepared_icons = self.icon_selector.prepare_icons_for_slide(icon_selections)

            # Insert text content first
            for placeholder_name, text_content in content.items():
                if placeholder_name in actual_placeholders:
                    placeholder_info = actual_placeholders[placeholder_name]

                    # Skip picture placeholders for now, handle them separately
                    if placeholder_info["type"] != PP_PLACEHOLDER.PICTURE:
                        placeholder_obj = placeholder_info["placeholder"]
                        self._set_placeholder_content(placeholder_obj, text_content)
                        print(f"  ✓ Set text content for '{placeholder_name}'")

            # Insert icons based on selections
            for icon_placeholder_name, icon_path in prepared_icons.items():
                if (
                    icon_placeholder_name in actual_placeholders
                    and icon_path
                    and os.path.exists(icon_path)
                ):

                    placeholder_info = actual_placeholders[icon_placeholder_name]
                    placeholder_obj = placeholder_info["placeholder"]

                    # Get the selected icon name
                    icon_name = icon_selections.get(icon_placeholder_name, "circle")

                    # Insert the icon
                    self._set_placeholder_content(placeholder_obj, icon_name)
                    print(
                        f"  ✓ Inserted icon '{icon_name}' for '{icon_placeholder_name}'"
                    )

            # Handle regular picture placeholders (not icons)
            regular_picture_placeholders = {
                name: info
                for name, info in actual_placeholders.items()
                if info["type"] == PP_PLACEHOLDER.PICTURE and "icon" not in name.lower()
            }

            for placeholder_name in regular_picture_placeholders:
                if placeholder_name in content:
                    # For regular picture placeholders, just leave them as placeholders
                    # The content is descriptive text, not an icon name
                    print(
                        f"  ✓ Skipped regular picture placeholder "
                        f"'{placeholder_name}' "
                        f"(content: '{content[placeholder_name][:50]}...')"
                    )
        else:
            # Regular slide without icons, use standard population
            self._populate_slide_with_actual_names(slide, content, actual_placeholders)

    def _populate_slide_with_actual_names(
        self,
        slide,
        content: Dict[str, str],
        actual_placeholders: Dict[str, Dict[str, Any]],
    ) -> None:
        """
        Populate slide using actual placeholder objects (no name matching needed)

        Args:
            slide: PowerPoint slide object
            content: Dictionary mapping placeholder names to content
            actual_placeholders: Dictionary of actual placeholder info
        """
        for placeholder_name, text_content in content.items():
            if placeholder_name in actual_placeholders:
                placeholder_obj = actual_placeholders[placeholder_name]["placeholder"]
                self._set_placeholder_content(placeholder_obj, text_content)
                print(f"  ✓ Set content for '{placeholder_name}'")
            else:
                print(
                    f"  ✗ Warning: Placeholder '{placeholder_name}' "
                    "not found in actual placeholders"
                )

    def _create_powerpoint_presentation(
        self, slide_contents: List[SlideContent], generated_images: dict = None
    ) -> Any:
        """
        Create the actual PowerPoint presentation

        Args:
            slide_contents: List of slide content objects
            generated_images: Dictionary of generated images by slide index

        Returns:
            PowerPoint presentation object
        """
        # Load template
        presentation = Presentation(self.template_path)

        # Create slides for each content without clearing existing slides
        # This avoids XML manipulation issues
        for i, slide_content in enumerate(slide_contents):
            self._add_slide_to_presentation(presentation, slide_content, generated_images, i)

        return presentation

    def _create_icon_aware_presentation(
        self, slide_contents: List[SlideContent], topic: str
    ) -> Any:
        """
        Create PowerPoint presentation with full icon management
        and smart content mapping

        Args:
            slide_contents: List of slide content objects
            topic: Presentation topic for icon selection context

        Returns:
            PowerPoint presentation object with proper icon handling
        """
        # Load template
        presentation = Presentation(self.template_path)

        # Create slides with icon-aware content population
        for slide_content in slide_contents:
            self._add_icon_aware_slide_to_presentation(
                presentation, slide_content, topic
            )

        return presentation

    def _add_icon_aware_slide_to_presentation(
        self, presentation: Any, slide_content: SlideContent, topic: str
    ) -> None:
        """
        Add a single slide to the presentation with full icon support using SAME APPROACH as individual slides

        Args:
            presentation: PowerPoint presentation object
            slide_content: Content for this slide
            topic: Presentation topic for icon context
        """
        # Get the layout
        layout = presentation.slide_layouts[slide_content.layout_index]

        # Add slide with the specified layout
        slide = presentation.slides.add_slide(layout)

        if slide_content.content:
            # CRITICAL FIX: Use EXACT same approach as individual slide generation for icon-aware slides too
            from .individual_slide_generator import IndividualSlideGenerator
            
            # Create temporary IndividualSlideGenerator instance to use its PROVEN methods
            individual_generator = IndividualSlideGenerator()
            
            # Apply content using the EXACT same method that works for individual slides
            individual_generator._apply_content_to_slide(
                slide,
                slide_content,
                self.template_path,
                getattr(self.content_generator, 'layouts_info', None),
                None,  # dynamic_models
                getattr(slide_content, 'html_image_path', None)
            )
            
            print(f"✅ Applied icon-aware content using PROVEN individual slide method")
        else:
            # No content provided but slide has placeholders
            print(f"  → Icon-aware slide created with no content")

    def _add_slide_to_presentation(
        self, presentation: Any, slide_content: SlideContent, generated_images: dict = None, slide_index: int = 0
    ) -> None:
        """
        Add a single slide to the presentation using EXACT SAME APPROACH as individual slides

        Args:
            presentation: PowerPoint presentation object
            slide_content: Content for this slide
            generated_images: Dictionary of generated images by slide index
            slide_index: Index of the current slide
        """
        # Get the layout
        layout = presentation.slide_layouts[slide_content.layout_index]

        # Add slide with the specified layout
        slide = presentation.slides.add_slide(layout)

        # CRITICAL FIX: Use EXACT same approach as individual slide generation
        if slide_content.content:
            # Import here to avoid circular imports (same as individual slides)
            from .individual_slide_generator import IndividualSlideGenerator
            
            # Create temporary IndividualSlideGenerator instance to use its PROVEN methods
            individual_generator = IndividualSlideGenerator()
            
            # Apply content using the EXACT same method that works for individual slides
            individual_generator._apply_content_to_slide(
                slide,
                slide_content,
                self.template_path,
                getattr(self.content_generator, 'layouts_info', None),
                None,  # dynamic_models
                getattr(slide_content, 'html_image_path', None)
            )
            
            print(f"✅ Applied content using PROVEN individual slide method")
        else:
            # No content provided but slide has placeholders
            print(f"  → Slide created with no content for slide {slide_index + 1}")
        
        # Handle generated images if available (using same approach as HTML method)
        if generated_images is not None and slide_index in generated_images:
            image_info = generated_images[slide_index]
            image_path = image_info.get("image_path")
            
            if image_path:
                # Use same approach as HTML method: check all picture placeholders on this slide
                for placeholder in slide.placeholders:
                    if (
                        hasattr(placeholder, "placeholder_format")
                        and placeholder.placeholder_format.type == PP_PLACEHOLDER.PICTURE
                    ):
                        # Get placeholder name (same as HTML method)
                        placeholder_name = getattr(placeholder, "name", "")
                        
                        # For generated images, we want Picture placeholders (not icons or HTML)
                        if "picture" in placeholder_name.lower() and "icon" not in placeholder_name.lower():
                            try:
                                # Set proper DPI metadata on the image to ensure correct sizing in PowerPoint
                                # This matches the HTML image insertion approach
                                try:
                                    from PIL import Image
                                    # Open the image and save with explicit DPI
                                    img = Image.open(image_path)
                                    # Generated images should use standard 96 DPI (they're already correct size)
                                    img.save(image_path, dpi=(96, 96))
                                    print(f"  - Set image DPI to 96 for proper PowerPoint display")
                                except Exception as e:
                                    print(f"  - Warning: Could not set DPI metadata: {e}")
                                    
                                self._replace_placeholder_with_image(
                                    placeholder, image_path, f"Generated Image {slide_index + 1}"
                                )
                                print(f"✅ Inserted generated image for slide {slide_index + 1}")
                                break
                            except Exception as e:
                                print(f"❌ Failed to insert generated image: {e}")
                else:
                    print(f"⚠️ No suitable picture placeholder found for generated image on slide {slide_index + 1}")

    def _get_actual_placeholder_info(self, slide) -> Dict[str, Dict[str, Any]]:
        """
        Get information about actual placeholders on a created slide

        Args:
            slide: PowerPoint slide object

        Returns:
            Dictionary mapping placeholder names to their info
        """
        placeholder_info = {}
        for placeholder in slide.placeholders:
            name = (
                placeholder.name or f"Placeholder_{placeholder.placeholder_format.idx}"
            )
            placeholder_info[name] = {
                "placeholder": placeholder,
                "type": placeholder.placeholder_format.type,
                "index": placeholder.placeholder_format.idx,
                "shape_type": placeholder.shape_type,
            }
        return placeholder_info

    def _get_layout_placeholder_formatting(self, slide, placeholder_idx: int) -> dict:
        """
        Get formatting from the layout placeholder
        
        Args:
            slide: PowerPoint slide object
            placeholder_idx: Index of the placeholder
            
        Returns:
            Dictionary with formatting properties from layout
        """
        formatting = {}
        try:
            layout = slide.slide_layout
            for layout_ph in layout.placeholders:
                if layout_ph.placeholder_format.idx == placeholder_idx:
                    if hasattr(layout_ph, 'text_frame') and layout_ph.text_frame:
                        text_frame = layout_ph.text_frame
                        if text_frame.paragraphs:
                            first_para = text_frame.paragraphs[0]
                            # Try to get font from runs first (most specific)
                            if first_para.runs:
                                font = first_para.runs[0].font
                                formatting['font_name'] = font.name
                                formatting['font_size'] = font.size
                            # Fall back to paragraph font
                            elif hasattr(first_para, 'font'):
                                font = first_para.font
                                formatting['font_name'] = font.name
                                formatting['font_size'] = font.size
                    break
        except Exception as e:
            print(f"Could not get layout formatting: {e}")
        return formatting
    
    def _capture_original_placeholders(self, slide) -> None:
        """
        Capture all original placeholders and their properties before modification
        
        Args:
            slide: PowerPoint slide object
        """
        self._original_placeholders = {}
        
        for placeholder in slide.placeholders:
            idx = placeholder.placeholder_format.idx
            
            # Capture placeholder properties
            properties = {
                'name': placeholder.name,
                'left': placeholder.left,
                'top': placeholder.top,
                'width': placeholder.width,
                'height': placeholder.height,
                'placeholder_format_type': placeholder.placeholder_format.type,
                'z_order_position': None  # Will be set based on position in shapes
            }
            
            # Find z-order position
            for i, shape in enumerate(slide.shapes):
                if shape == placeholder:
                    properties['z_order_position'] = i
                    break
            
            # Capture text properties if it's a text placeholder
            if hasattr(placeholder, 'text_frame'):
                text_frame = placeholder.text_frame
                properties['text'] = text_frame.text if text_frame else ""
                
                # Capture paragraph and font properties from first paragraph
                if text_frame and text_frame.paragraphs:
                    first_para = text_frame.paragraphs[0]
                    properties['paragraph_alignment'] = first_para.alignment
                    properties['paragraph_level'] = first_para.level
                    
                    # Capture font properties from first run if exists
                    if first_para.runs:
                        first_run = first_para.runs[0]
                        font = first_run.font
                        properties['font_name'] = font.name
                        properties['font_size'] = font.size
                        properties['font_bold'] = font.bold
                        properties['font_italic'] = font.italic
                        # Safely get font color
                        try:
                            properties['font_color'] = font.color.rgb if font.color else None
                        except AttributeError:
                            properties['font_color'] = None
            
            self._original_placeholders[idx] = properties
    
    def _restore_unused_placeholders(self, slide, content: Dict[str, str], layout_info: Dict = None) -> None:
        """
        Restore placeholders that weren't filled with content as text boxes with original styling
        
        Args:
            slide: PowerPoint slide object
            content: Dictionary of content that was applied
            layout_info: Optional layout information with custom names
        """
        if not hasattr(self, '_original_placeholders'):
            return
        
        # Determine which placeholders were used
        used_placeholder_names = set()
        
        # Add content keys to used names
        used_placeholder_names.update(content.keys())
        
        # Get mapping of idx to custom names if available
        idx_to_custom_name = {}
        if layout_info:
            for placeholder_info in layout_info.get("placeholders", []):
                idx = placeholder_info.get("index")
                custom_name = placeholder_info.get("name")
                if idx is not None and custom_name:
                    idx_to_custom_name[idx] = custom_name
        
        # Check each original placeholder
        for idx, properties in self._original_placeholders.items():
            custom_name = idx_to_custom_name.get(idx, properties['name'])
            
            # Skip if this placeholder was used
            if custom_name in used_placeholder_names:
                continue
            
            # Skip if it's a picture placeholder that was replaced
            if properties['placeholder_format_type'] == 18:  # PICTURE type
                # Check if the placeholder still exists (wasn't replaced)
                placeholder_exists = False
                for placeholder in slide.placeholders:
                    if placeholder.placeholder_format.idx == idx:
                        placeholder_exists = True
                        break
                
                if not placeholder_exists:
                    continue  # Was replaced with an image
            
            # Find if placeholder still exists
            current_placeholder = None
            for placeholder in slide.placeholders:
                if placeholder.placeholder_format.idx == idx:
                    current_placeholder = placeholder
                    break
            
            # If placeholder exists and has no new content, ensure it keeps original text
            if current_placeholder and hasattr(current_placeholder, 'text_frame'):
                if properties.get('text'):
                    # Restore original text if it was cleared
                    if not current_placeholder.text_frame.text:
                        current_placeholder.text_frame.text = properties['text']
    
    def _populate_slide_placeholders(self, slide, content: Dict[str, str]) -> None:
        """
        Populate slide placeholders with content using intelligent mapping
        Also preserves unused placeholders with their original styling

        Args:
            slide: PowerPoint slide object
            content: Dictionary mapping placeholder names to content
        """
        # First, capture all original placeholders and their properties
        self._capture_original_placeholders(slide)
        
        # Auto-add LOCKED_ backgrounds from template folder
        from pathlib import Path
        template_folder = Path(self.template_path).parent
        if template_folder.exists():
            # Check the slide layout for LOCKED_ placeholders
            slide_layout = slide.slide_layout
            for layout_ph in slide_layout.placeholders:
                ph_name = layout_ph.name
                if ph_name and ph_name.startswith("LOCKED_"):
                    # Check if we have a PNG file for this LOCKED_ placeholder
                    png_path = template_folder / f"{ph_name}.png"
                    if png_path.exists() and ph_name not in content:
                        # Auto-add the LOCKED_ background to content
                        content[ph_name] = str(png_path)
                        print(f"  🔒 Auto-adding LOCKED_ background: {ph_name}")
        
        
        # Get layout information for this slide to access custom placeholder mapping
        slide_layout_index = None
        for i, layout in enumerate(slide.slide_layout.slide_master.slide_layouts):
            if layout == slide.slide_layout:
                slide_layout_index = i
                break

        if slide_layout_index is not None and hasattr(self, "content_generator"):
            # Get layout info which has the custom name to index mapping
            layout_info = self.content_generator.layouts_info.get(slide_layout_index)
            if layout_info:
                self._populate_with_layout_mapping(slide, content, layout_info)
                # Restore any unused placeholders
                self._restore_unused_placeholders(slide, content, layout_info)
                return

        # Fallback to original method if no layout mapping available
        self._populate_with_name_matching(slide, content)
        # Restore any unused placeholders
        self._restore_unused_placeholders(slide, content)

    def _populate_with_layout_mapping(
        self, slide, content: Dict[str, str], layout_info: Dict
    ) -> None:
        """
        Populate placeholders using layout analysis mapping (handles custom names)
        MAINTAINS EXACT Z-ORDER from template

        Args:
            slide: PowerPoint slide object
            content: Dictionary mapping custom placeholder names to content
            layout_info: Layout information with custom name to index mapping
        """
        # Create mapping from placeholder indices to placeholder objects
        placeholder_by_index = {}
        for placeholder in slide.placeholders:
            idx = placeholder.placeholder_format.idx
            placeholder_by_index[idx] = placeholder

        print("  → Mapping content using layout analysis:")
        
        # Track original z-order positions before any replacements
        original_z_order = {}
        for i, shape in enumerate(slide.shapes):
            if hasattr(shape, 'placeholder_format'):
                original_z_order[shape.placeholder_format.idx] = i
        
        # Track which placeholders will be replaced with images
        replaced_placeholders = {}

        # Process ALL placeholders in any order (we'll fix z-order after)
        for layout_placeholder in layout_info.get("placeholders", []):
            custom_name = layout_placeholder["name"]
            placeholder_index = layout_placeholder["index"]
            
            if custom_name in content:
                if placeholder_index in placeholder_by_index:
                    placeholder_obj = placeholder_by_index[placeholder_index]
                    text_content = content[custom_name]
                    
                    # Track if this will be replaced with an image
                    is_image = (hasattr(placeholder_obj, 'placeholder_format') and 
                               placeholder_obj.placeholder_format.type == 18 and  # PICTURE type
                               (custom_name.startswith("LOCKED_") or 
                                self._is_image_path(text_content)))
                    
                    if is_image:
                        replaced_placeholders[placeholder_index] = original_z_order.get(placeholder_index, 0)
                    
                    # Pass custom name info to determine if this is an icon placeholder
                    self._set_placeholder_content_with_custom_name(
                        placeholder_obj, text_content, custom_name
                    )
                    
                    if custom_name.startswith("LOCKED_"):
                        print(f"    🔒 '{custom_name}' → Index {placeholder_index} (z-order: {original_z_order.get(placeholder_index, 'unknown')})")
                    else:
                        print(f"    ✓ '{custom_name}' → Index {placeholder_index}")
                else:
                    print(f"    ✗ Index {placeholder_index} not found in slide for '{custom_name}'")
            else:
                print(f"    - No content for '{custom_name}'")
        
        # After all replacements, fix z-order to match template
        self._restore_z_order(slide, original_z_order, replaced_placeholders)

    def _restore_z_order(self, slide, original_z_order: dict, replaced_placeholders: dict) -> None:
        """
        Restore z-order to match the original template order
        
        Args:
            slide: PowerPoint slide object
            original_z_order: Original z-order positions from template
            replaced_placeholders: Dict of placeholder indices that were replaced with images
        """
        try:
            # For each replaced placeholder, find its image and ensure correct position
            for placeholder_idx, target_position in sorted(replaced_placeholders.items(), key=lambda x: x[1]):
                # The image that replaced this placeholder should be at target_position
                # Find all images
                images = [shape for shape in slide.shapes if hasattr(shape, 'image')]
                
                # Try to identify which image replaced this placeholder
                # Images are usually added in order, so we can match by position
                for img_shape in images:
                    current_position = list(slide.shapes).index(img_shape)
                    
                    # Check if this image is close to where we expect it
                    # (within reasonable range considering other replacements)
                    if abs(current_position - target_position) <= len(images):
                        # This is likely our replacement image
                        if current_position != target_position:
                            # Need to move it
                            shape_element = img_shape._element
                            parent = shape_element.getparent()
                            
                            # Remove from current position
                            parent.remove(shape_element)
                            
                            # Insert at target position
                            # Get all current children for reference
                            children = list(parent)
                            
                            if target_position < len(children):
                                # Insert before the element at target position
                                parent.insert(target_position, shape_element)
                            else:
                                # Append at the end
                                parent.append(shape_element)
                            
                            print(f"  → Moved {img_shape.name} from position {current_position} to {target_position}")
                        break
                        
        except Exception as e:
            print(f"  ⚠️ Warning: Could not fully restore z-order: {e}")

    def _populate_with_name_matching(self, slide, content: Dict[str, str]) -> None:
        """
        Fallback method using name matching with layout placeholder names
        Args:
            slide: PowerPoint slide object
            content: Dictionary mapping placeholder names to content
        """
        # Create mapping from layout placeholder names to slide placeholders by index
        placeholder_map = {}
        layout_placeholder_map = {}
        
        # Get the slide's layout
        slide_layout = slide.slide_layout
        
        # Build a map of layout placeholder names by index
        for layout_ph in slide_layout.placeholders:
            idx = layout_ph.placeholder_format.idx
            layout_placeholder_map[idx] = layout_ph.name
        
        # Map slide placeholders using their indices to get layout names
        for placeholder in slide.placeholders:
            idx = placeholder.placeholder_format.idx
            
            # Use layout name if available, otherwise use slide placeholder name
            if idx in layout_placeholder_map:
                name = layout_placeholder_map[idx]
                print(f"  📍 Mapped placeholder idx {idx}: '{name}'")
            else:
                name = placeholder.name or f"Placeholder_{idx}"
            
            placeholder_map[name] = placeholder
        
        # Debug: Show available placeholder names
        print(f"  🔍 Available placeholders for content mapping:")
        for name in placeholder_map.keys():
            ph = placeholder_map[name]
            ph_type = "PICTURE" if ph.placeholder_format.type == 18 else "TEXT"
            print(f"     • '{name}' ({ph_type})")
        
        # Fill placeholders with content
        for placeholder_name, text_content in content.items():
            # Don't skip LOCKED_ placeholders if they have content (they should have image paths)
            if placeholder_name.startswith("LOCKED_") and not text_content:
                print(f"  - No content for locked placeholder '{placeholder_name}'")
                continue
                
            if placeholder_name in placeholder_map:
                placeholder = placeholder_map[placeholder_name]
                print(f"  ✅ Found placeholder '{placeholder_name}'")
                self._set_placeholder_content(placeholder, text_content)
            else:
                # Try to find placeholder by partial name match
                matched_placeholder = self._find_placeholder_by_partial_match(
                    placeholder_map, placeholder_name
                )
                if matched_placeholder:
                    print(f"  ✅ Found placeholder '{placeholder_name}' by partial match")
                    self._set_placeholder_content(matched_placeholder, text_content)
                else:
                    print(f"  ⚠️ Warning: Placeholder '{placeholder_name}' not found in mapping")

    def _find_placeholder_by_partial_match(
        self, placeholder_map: Dict[str, Any], target_name: str
    ) -> Optional[Any]:
        """
        Find placeholder by partial name matching

        Args:
            placeholder_map: Dictionary of placeholder names to objects
            target_name: Target placeholder name to find

        Returns:
            Placeholder object if found, None otherwise
        """
        target_lower = target_name.lower()

        # Try exact match first
        if target_name in placeholder_map:
            return placeholder_map[target_name]

        # Try case-insensitive match
        for name, placeholder in placeholder_map.items():
            if name.lower() == target_lower:
                return placeholder

        # Enhanced matching for common placeholder name patterns
        # Handle PowerPoint's automatic naming patterns
        for name, placeholder in placeholder_map.items():
            name_lower = name.lower()

            # Handle "Title" matching "Title 1"
            if target_lower == "title" and name_lower.startswith("title"):
                return placeholder

            # Handle "Subtitle" matching "Text Placeholder 2" (common for subtitle)
            if target_lower == "subtitle" and "text placeholder" in name_lower:
                return placeholder

            # Handle "Presenter" matching remaining text placeholders
            if target_lower == "presenter" and "text placeholder" in name_lower:
                # Try to find the last text placeholder for presenter
                continue

            # Handle content placeholders
            if (
                "content placeholder" in target_lower
                and "content placeholder" in name_lower
            ):
                return placeholder

            # Handle "Text Content" matching "Content Placeholder"
            if target_lower == "text content" and "content placeholder" in name_lower:
                return placeholder

            # Handle chart placeholders
            if (
                "chart placeholder" in target_lower
                and "chart placeholder" in name_lower
            ):
                return placeholder

            # Handle picture placeholders
            if "picture" in target_lower and "picture" in name_lower:
                return placeholder

            # Generic partial matching for other cases
            if target_lower in name_lower or name_lower in target_lower:
                return placeholder

        # Special handling for "Presenter" - use the last available text placeholder
        if target_lower == "presenter":
            text_placeholders = [
                (name, placeholder)
                for name, placeholder in placeholder_map.items()
                if "text placeholder" in name.lower()
            ]
            if text_placeholders:
                # Sort by name and take the last one (highest number)
                text_placeholders.sort(key=lambda x: x[0])
                return text_placeholders[-1][1]

        return None

    def _set_placeholder_content_with_custom_name(
        self, placeholder, content: str, custom_name: str
    ) -> None:
        """
        Set content for a placeholder using custom name to determine handling

        Args:
            placeholder: PowerPoint placeholder object
            content: Text content to set, chart data, icon name, or HTML content
            custom_name: Custom placeholder name from template analysis
        """
        try:
            # Check if this is a picture placeholder
            if (
                hasattr(placeholder, "placeholder_format")
                and placeholder.placeholder_format.type == PP_PLACEHOLDER.PICTURE
            ):
                # Use custom name to determine if this is an icon placeholder
                if "icon" in custom_name.lower():
                    # Handle as icon placeholder
                    self._insert_icon_into_placeholder(placeholder, content)
                elif self._is_html_content(content) and self.html_renderer:
                    # Handle as HTML visualization
                    self._insert_html_visualization_into_placeholder(
                        placeholder, content, custom_name
                    )
                elif self._is_image_path(content):
                    # Handle as image file path
                    self._insert_image_into_placeholder(placeholder, content)
                else:
                    # Handle as regular picture placeholder - leave as placeholder
                    print(
                        f"  ✓ Left regular picture placeholder '{custom_name}' "
                        f"as placeholder (content: '{content[:30]}...')"
                    )
                return

            # Check if this is a chart placeholder
            if (
                hasattr(placeholder, "placeholder_format")
                and placeholder.placeholder_format.type == PP_PLACEHOLDER.CHART
            ):
                # Try to create chart from content
                if self._create_chart_from_content(placeholder, content):
                    return  # Chart created successfully
                # If chart creation fails, fall through to text insertion

            if hasattr(placeholder, "text_frame"):
                # Text placeholder - preserve original formatting from layout
                # Get the placeholder's index to find layout formatting
                placeholder_idx = placeholder.placeholder_format.idx
                layout_formatting = self._get_layout_placeholder_formatting(
                    placeholder.part.slide, placeholder_idx
                )
                
                # Set text while preserving formatting
                self._set_text_preserving_formatting_with_layout(
                    placeholder, content, layout_formatting
                )
            elif hasattr(placeholder, "text"):
                # Simple text placeholder
                placeholder.text = content
            else:
                print(
                    f"Warning: Unknown placeholder type for content: {content[:50]}..."
                )
        except Exception as e:
            print(f"Error setting placeholder content: {e}")

    def _remove_empty_picture_placeholder(self, placeholder):
        """
        Remove an empty picture placeholder from the slide to prevent corruption
        
        Args:
            placeholder: The placeholder to remove
        """
        try:
            # Get the slide from the placeholder
            slide = placeholder.part.slide
            
            # Find and remove the placeholder shape from the slide
            for shape in slide.shapes:
                if shape == placeholder:
                    # Remove the shape element from the slide's shape tree
                    slide.shapes._spTree.remove(shape._element)
                    print(f"    → Removed empty picture placeholder from slide")
                    break
        except Exception as e:
            print(f"    → Warning: Could not remove empty placeholder: {e}")
            # If removal fails, at least try to make it invisible
            try:
                placeholder.width = 0
                placeholder.height = 0
            except:
                pass

    def _is_html_content(self, content: str) -> bool:
        """
        Check if content contains HTML that should be rendered as visualization

        Args:
            content: Content string to check

        Returns:
            True if content appears to be HTML for visualization
        """
        html_indicators = [
            "timeline:",
            "process:",
            "flowchart:",
            "diagram:",
            "<html>",
            "<!DOCTYPE",
            "<div",
            "<timeline>",
            "<process>",
        ]
        content_lower = content.lower().strip()
        return any(indicator in content_lower for indicator in html_indicators)

    def _is_image_path(self, content: str) -> bool:
        """
        Check if content is a file path pointing to an image file
        
        Args:
            content: Content string to check
            
        Returns:
            True if content appears to be an image file path
        """
        if not content or not isinstance(content, str):
            return False
            
        content = content.strip()
        
        # Check for common image file extensions
        image_extensions = ['.png', '.jpg', '.jpeg', '.gif', '.bmp', '.tiff', '.webp']
        content_lower = content.lower()
        
        # Check if it ends with an image extension
        has_image_extension = any(content_lower.endswith(ext) for ext in image_extensions)
        
        # Check if it looks like a file path (contains directory separators)
        looks_like_path = ('/' in content or '\\' in content or content.startswith('.'))
        
        return has_image_extension and looks_like_path

    def _insert_html_visualization_into_placeholder(
        self, placeholder, content: str, custom_name: str
    ) -> None:
        """
        Render HTML content as an image and insert into picture placeholder

        Args:
            placeholder: PowerPoint picture placeholder object
            content: HTML content or visualization instructions
            custom_name: Custom placeholder name for context
        """
        try:
            if not self.html_renderer:
                print(f"HTML renderer not available for '{custom_name}'")
                return

            # Parse content to determine visualization type
            html_content = self._generate_html_from_content(content)

            if not html_content:
                print(f"Could not generate HTML for content: {content[:50]}...")
                return

            # Create temporary file for rendered image
            import tempfile

            with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as temp_file:
                temp_image_path = temp_file.name

            # Get placeholder dimensions for dynamic sizing
            # PowerPoint internally uses EMUs (English Metric Units)
            # 1 inch = 914400 EMUs, and we need to convert to pixels
            # Using standard screen resolution of 96 DPI for consistency
            DPI = 96
            EMU_PER_INCH = 914400
            EMU_PER_PIXEL = EMU_PER_INCH / DPI  # 9525 EMUs per pixel at 96 DPI
            
            placeholder_width_px = int(placeholder.width.emu / EMU_PER_PIXEL)
            placeholder_height_px = int(placeholder.height.emu / EMU_PER_PIXEL)
            
            print(f"  - Placeholder dimensions: {placeholder_width_px}x{placeholder_height_px}px")
            
            # Render HTML to image with 2x resolution for high quality
            # The HTML renderer internally uses device scale factor for crisp rendering
            success = self.html_renderer.render_html_to_image(
                html_content=html_content,
                output_path=temp_image_path,
                width=placeholder_width_px * 2,  # 2x resolution for high quality
                height=placeholder_height_px * 2,  # 2x resolution for high quality
            )

            if success and os.path.exists(temp_image_path):
                # Set proper DPI metadata on the image to ensure correct sizing in PowerPoint
                try:
                    from PIL import Image
                    # Open the image and save with explicit DPI
                    img = Image.open(temp_image_path)
                    # Save with 192 DPI (2x of 96) since we rendered at 2x resolution
                    img.save(temp_image_path, dpi=(192, 192))
                    print(f"  - Set image DPI to 192 for proper PowerPoint display")
                except Exception as e:
                    print(f"  - Warning: Could not set DPI metadata: {e}")
                
                # Replace placeholder with rendered image
                self._replace_placeholder_with_image(
                    placeholder, temp_image_path, custom_name
                )

                # Clean up temporary file
                try:
                    os.unlink(temp_image_path)
                except OSError:
                    pass  # Ignore cleanup errors

                print(f"✅ Inserted HTML visualization for '{custom_name}'")
            else:
                print(f"Failed to render HTML visualization for '{custom_name}'")

        except Exception as e:
            print(f"Error inserting HTML visualization '{custom_name}': {e}")

    def _generate_html_from_content(self, content: str) -> Optional[str]:
        """
        Generate HTML from content string based on visualization type

        Args:
            content: Content string with visualization instructions

        Returns:
            HTML string or None if can't generate
        """
        content_lower = content.lower().strip()

        # Timeline visualization
        if "timeline:" in content_lower:
            return self._create_timeline_from_content(content)

        # Process flow visualization
        if "process:" in content_lower:
            return self._create_process_flow_from_content(content)

        # Direct HTML content
        if any(tag in content_lower for tag in ["<html>", "<!doctype", "<div"]):
            return content

        # Default: treat as timeline if it has timeline-like structure
        if self._looks_like_timeline(content):
            return self._create_timeline_from_content(content)

        return None

    def _create_timeline_from_content(self, content: str) -> str:
        """Create timeline HTML from content description"""
        # Parse timeline events from content
        events = self._parse_timeline_events(content)

        # Extract title from content
        lines = content.split("\n")
        title = "Timeline"
        for line in lines:
            if line.strip() and not line.lower().startswith("timeline:"):
                title = line.strip()
                break

        if self.html_renderer:
            return self.html_renderer.create_timeline_html(
                events=events, title=title, theme="ekona"
            )
        return ""

    def _create_process_flow_from_content(self, content: str) -> str:
        """Create process flow HTML from content description"""
        # Parse process steps from content
        steps = self._parse_process_steps(content)

        # Extract title from content
        lines = content.split("\n")
        title = "Process Flow"
        for line in lines:
            if line.strip() and not line.lower().startswith("process:"):
                title = line.strip()
                break

        if self.html_renderer:
            return self.html_renderer.create_process_flow_html(
                steps=steps, title=title, theme="ekona"
            )
        return ""

    def _parse_timeline_events(self, content: str) -> List[Dict[str, str]]:
        """Parse timeline events from content string"""
        events = []
        lines = content.split("\n")

        current_event = {}
        for line in lines:
            line = line.strip()
            if not line or line.lower().startswith("timeline:"):
                continue

            # Look for date patterns
            if any(char.isdigit() for char in line) and len(line) < 20:
                if current_event:
                    events.append(current_event)
                current_event = {"date": line, "title": "", "description": ""}
            elif not current_event.get("title"):
                current_event["title"] = line
            else:
                current_event["description"] = line

        if current_event:
            events.append(current_event)

        # Default events if parsing fails
        if not events:
            events = [
                {
                    "date": "2024 Q1",
                    "title": "Project Start",
                    "description": "Initiative launched",
                },
                {
                    "date": "2024 Q2",
                    "title": "Development",
                    "description": "Core development phase",
                },
                {
                    "date": "2024 Q3",
                    "title": "Testing",
                    "description": "Quality assurance phase",
                },
                {"date": "2024 Q4", "title": "Launch", "description": "Product launch"},
            ]

        return events

    def _parse_process_steps(self, content: str) -> List[Dict[str, str]]:
        """Parse process steps from content string"""
        steps = []
        lines = content.split("\n")

        for line in lines:
            line = line.strip()
            if not line or line.lower().startswith("process:"):
                continue

            # Simple parsing - each line is a step
            if line:
                steps.append({"title": line, "description": f"Complete {line.lower()}"})

        # Default steps if parsing fails
        if not steps:
            steps = [
                {"title": "Plan", "description": "Define requirements and strategy"},
                {"title": "Design", "description": "Create detailed design"},
                {"title": "Build", "description": "Implement solution"},
                {"title": "Test", "description": "Verify quality and functionality"},
                {"title": "Deploy", "description": "Launch to production"},
            ]

        return steps

    def _looks_like_timeline(self, content: str) -> bool:
        """Check if content looks like timeline data"""
        timeline_keywords = [
            "date",
            "time",
            "year",
            "month",
            "quarter",
            "phase",
            "milestone",
            "event",
            "history",
            "chronology",
        ]
        content_lower = content.lower()
        return any(keyword in content_lower for keyword in timeline_keywords)

    def _replace_placeholder_with_image(
        self, placeholder, image_path: str, name: str
    ) -> None:
        """
        Replace a placeholder with an image file while maintaining z-order

        Args:
            placeholder: PowerPoint placeholder object
            image_path: Path to image file
            name: Name for the new image shape
        """
        try:
            # Get placeholder properties before replacement
            left = placeholder.left
            top = placeholder.top
            width = placeholder.width
            height = placeholder.height
            
            # Store the placeholder's z-order position
            z_order_position = None

            # Get the slide and shapes collection
            slide = placeholder.part.slide
            shapes = slide.shapes

            # Find the placeholder's position in shapes for z-order
            placeholder_idx = None
            for i, shape in enumerate(shapes):
                if shape == placeholder:
                    placeholder_idx = i
                    z_order_position = i
                    break

            if placeholder_idx is not None:
                # Store reference to the element's parent and next sibling for z-order preservation
                placeholder_element = placeholder._element
                parent = placeholder_element.getparent()
                next_sibling = placeholder_element.getnext()
                
                # Delete the placeholder
                shapes._spTree.remove(placeholder._element)

                # Add the image with DPI-aware sizing
                # First, let PowerPoint auto-size based on image DPI, then adjust if needed
                picture = shapes.add_picture(image_path, left, top)
                
                # Check if auto-sizing worked correctly
                auto_width = picture.width
                auto_height = picture.height
                
                # Compare auto-sized dimensions with placeholder dimensions
                EMU_PER_INCH = 914400
                tolerance = EMU_PER_INCH * 0.2  # Increased to 0.2 inch tolerance (~19px)
                
                width_diff = abs(auto_width.emu - width.emu)
                height_diff = abs(auto_height.emu - height.emu)
                
                auto_width_px = auto_width.emu // 9525
                auto_height_px = auto_height.emu // 9525
                target_width_px = width.emu // 9525
                target_height_px = height.emu // 9525
                
                print(f"  - Auto-sized to: {auto_width_px}x{auto_height_px}px")
                print(f"  - Placeholder: {target_width_px}x{target_height_px}px")
                
                if width_diff > tolerance or height_diff > tolerance:
                    print(f"  - Size difference detected, applying crop-to-fill")
                    
                    # Crop to fill: Scale image to cover entire placeholder, then crop excess
                    auto_aspect = auto_width_px / auto_height_px if auto_height_px > 0 else 1
                    target_aspect = target_width_px / target_height_px if target_height_px > 0 else 1
                    
                    if abs(auto_aspect - target_aspect) < 0.01:  # Aspect ratios match
                        # Just use placeholder dimensions since aspect ratios match
                        picture.width = width
                        picture.height = height
                        print(f"  - Aspect ratios match, using exact placeholder dimensions")
                    else:
                        # Use PowerPoint's crop feature to fill the placeholder
                        # Set the image to placeholder size (this will distort temporarily)
                        picture.width = width
                        picture.height = height
                        
                        # Apply crop to maintain aspect ratio
                        # Calculate how much to scale to fill (use max instead of min)
                        scale_width = target_width_px / auto_width_px
                        scale_height = target_height_px / auto_height_px
                        scale = max(scale_width, scale_height)  # Scale to fill (not fit)
                        
                        # The image is now sized to the placeholder
                        # PowerPoint will automatically center the image content
                        print(f"  - Applied crop-to-fill at placeholder size: {target_width_px}x{target_height_px}px")
                        
                        # Set crop properties to center the image
                        try:
                            # Access the crop properties
                            picture.crop_left = 0
                            picture.crop_right = 0
                            picture.crop_top = 0
                            picture.crop_bottom = 0
                            
                            # Calculate crop amounts if aspect ratios don't match
                            if auto_aspect > target_aspect:
                                # Image is wider - crop left and right
                                crop_amount = (1 - (target_aspect / auto_aspect)) / 2
                                picture.crop_left = crop_amount
                                picture.crop_right = crop_amount
                                print(f"  - Cropping sides by {crop_amount:.1%} each")
                            else:
                                # Image is taller - crop top and bottom
                                crop_amount = (1 - (auto_aspect / target_aspect)) / 2
                                picture.crop_top = crop_amount
                                picture.crop_bottom = crop_amount
                                print(f"  - Cropping top/bottom by {crop_amount:.1%} each")
                        except:
                            # If crop properties aren't available, the image will just be stretched
                            print(f"  - Using fill mode (image will fill placeholder)")
                else:
                    print(f"  - Auto-sizing worked correctly, keeping auto dimensions")
                
                picture.name = f"{name}_visualization"
                
                # Restore z-order position if we have it
                if z_order_position is not None and z_order_position < len(shapes):
                    # Move the picture to the original z-order position
                    picture_element = picture._element
                    
                    # Remove from current position
                    parent.remove(picture_element)
                    
                    # CRITICAL: Check if we're about to insert before grpSpPr
                    # The grpSpPr element must always come immediately after nvGrpSpPr
                    if next_sibling is not None:
                        # Check if next_sibling is grpSpPr
                        if next_sibling.tag.endswith('grpSpPr'):
                            # Find the element after grpSpPr to insert before that instead
                            grpSpPr_next = next_sibling.getnext()
                            if grpSpPr_next is not None:
                                parent.insert(parent.index(grpSpPr_next), picture_element)
                                print(f"  - Adjusted z-order to avoid breaking grpSpPr position")
                            else:
                                # grpSpPr was last, append after it
                                parent.append(picture_element)
                                print(f"  - Appended after grpSpPr to maintain structure")
                        else:
                            # Safe to insert at original position
                            parent.insert(parent.index(next_sibling), picture_element)
                            print(f"  - Maintained z-order position: {z_order_position}")
                    else:
                        # Was at the end, append
                        parent.append(picture_element)
                        print(f"  - Maintained z-order position: {z_order_position} (end)")

                print("✅ Replaced placeholder with visualization image")
            else:
                print("Warning: Could not find placeholder in shapes collection")

        except Exception as e:
            print(f"Error replacing placeholder with image: {e}")

    def _set_placeholder_content(self, placeholder, content) -> None:
        """
        Set content for a placeholder while preserving original formatting

        Args:
            placeholder: PowerPoint placeholder object
            content: Text content (str), list of items, chart data, icon name, or HTML content
        """
        try:
            # Handle list content first (for bullet points)
            if isinstance(content, list):
                if hasattr(placeholder, "text_frame"):
                    # Join list items with newlines for bullet points
                    bullet_text = "\n".join(str(item) for item in content)
                    # Get layout formatting
                    placeholder_idx = placeholder.placeholder_format.idx
                    layout_formatting = self._get_layout_placeholder_formatting(
                        placeholder.part.slide, placeholder_idx
                    )
                    # Set text with preserved formatting
                    self._set_text_preserving_formatting_with_layout(
                        placeholder, bullet_text, layout_formatting
                    )
                else:
                    print(f"Warning: Cannot set list content on non-text placeholder")
                return
            
            # Check if this is a picture placeholder
            if (
                hasattr(placeholder, "placeholder_format")
                and placeholder.placeholder_format.type == PP_PLACEHOLDER.PICTURE
            ):
                # Get placeholder name for detection logic
                placeholder_name = getattr(placeholder, "name", "")

                # Check for HTML visualization content
                if self._is_html_content(content) and self.html_renderer:
                    # Handle as HTML visualization
                    self._insert_html_visualization_into_placeholder(
                        placeholder, content, placeholder_name
                    )
                    return
                elif "icon" in placeholder_name.lower():
                    # Handle as icon placeholder
                    self._insert_icon_into_placeholder(placeholder, content)
                    return
                elif self._is_image_path(content):
                    # Handle as image file path
                    self._insert_image_into_placeholder(placeholder, content)
                    return
                else:
                    # Handle as regular picture placeholder - leave as placeholder
                    print(
                        f"  ✓ Left regular picture placeholder '{placeholder_name}' "
                        f"as placeholder (content: '{content[:30]}...')"
                    )
                return

            # Check if this is a chart placeholder
            if (
                hasattr(placeholder, "placeholder_format")
                and placeholder.placeholder_format.type == PP_PLACEHOLDER.CHART
            ):
                # Try to create chart from content
                if self._create_chart_from_content(placeholder, content):
                    return  # Chart created successfully
                # If chart creation fails, fall through to text insertion

            if hasattr(placeholder, "text_frame"):
                # Text placeholder - preserve original formatting from layout
                # Get the placeholder's index to find layout formatting
                placeholder_idx = placeholder.placeholder_format.idx
                layout_formatting = self._get_layout_placeholder_formatting(
                    placeholder.part.slide, placeholder_idx
                )
                
                # Set text while preserving formatting
                self._set_text_preserving_formatting_with_layout(
                    placeholder, content, layout_formatting
                )
            elif hasattr(placeholder, "text"):
                # Simple text placeholder
                placeholder.text = content
            else:
                print(
                    f"Warning: Unknown placeholder type for content: {content[:50]}..."
                )
        except Exception as e:
            print(f"Error setting placeholder content: {e}")

    def _insert_icon_into_placeholder(self, placeholder, icon_name: str) -> None:
        """
        Insert an icon into a picture placeholder by replacing
        the placeholder

        Args:
            placeholder: PowerPoint picture placeholder object
            icon_name: Name of the icon to insert
        """
        try:
            # Prepare the icon (convert to PNG if needed)
            icon_path = self.icon_manager.prepare_icon(icon_name, size=128)

            if not icon_path or not os.path.exists(icon_path):
                print(f"Warning: Icon '{icon_name}' not found or failed to prepare")
                return

            # Get placeholder properties before replacement
            left = placeholder.left
            top = placeholder.top
            width = placeholder.width
            height = placeholder.height
            name = getattr(placeholder, "name", f"Icon_{icon_name}")

            # Get the slide and shapes collection
            slide = placeholder.part.slide
            shapes = slide.shapes

            # Remove the original placeholder
            # First find the placeholder in the shapes collection
            placeholder_idx = None
            for i, shape in enumerate(shapes):
                if shape == placeholder:
                    placeholder_idx = i
                    break

            if placeholder_idx is not None:
                # Delete the placeholder
                shapes._spTree.remove(placeholder._element)

                # Add the icon image in the same position and size
                picture = shapes.add_picture(icon_path, left, top, width, height)

                # Set the picture name
                picture.name = f"{name}_icon"

                print(f"✅ Inserted icon '{icon_name}' into picture placeholder")
            else:
                print("Warning: Could not find placeholder in shapes collection")

        except Exception as e:
            print(f"Error inserting icon '{icon_name}': {e}")
            import traceback

            traceback.print_exc()

    def _insert_image_into_placeholder(self, placeholder, image_path: str) -> None:
        """
        Insert an image file into a picture placeholder by replacing the placeholder
        
        Args:
            placeholder: PowerPoint picture placeholder object
            image_path: Path to the image file to insert
        """
        try:
            # Check if image file exists
            if not os.path.exists(image_path):
                print(f"Warning: Image file '{image_path}' not found")
                return
            
            # Use the existing well-tested _replace_placeholder_with_image method
            placeholder_name = getattr(placeholder, 'name', 'Image')
            self._replace_placeholder_with_image(placeholder, image_path, f"{placeholder_name}_image")
            print(f"✅ Inserted image '{image_path}' into picture placeholder")
                
        except Exception as e:
            print(f"Error inserting image '{image_path}': {e}")
            import traceback
            traceback.print_exc()

    def _create_chart_from_content(self, placeholder, content: str) -> bool:
        """
        Attempt to create a chart from content string

        Args:
            placeholder: Chart placeholder object
            content: Content string that might contain chart instructions

        Returns:
            True if chart was created successfully, False otherwise
        """
        try:
            # For now, create a sample chart based on content keywords
            # TODO: Parse actual chart data from LLM in future enhancement
            chart_data = self._parse_chart_content(content)
            return self.chart_generator.create_chart(placeholder, chart_data)

        except Exception as e:
            print(f"Error creating chart: {e}")
            return False

    def _parse_chart_content(self, content: str) -> Dict[str, Any]:
        """
        Parse chart content and create appropriate chart data

        Args:
            content: Text content describing the chart

        Returns:
            Dictionary with chart data structure
        """
        # Simple parsing logic - can be enhanced with LLM chart data generation
        content_lower = content.lower()

        # Determine chart type from content
        chart_type = "column"  # default
        if any(word in content_lower for word in ["pie", "distribution", "percentage"]):
            chart_type = "pie"
        elif any(word in content_lower for word in ["line", "trend", "over time"]):
            chart_type = "line"
        elif any(word in content_lower for word in ["bar", "horizontal"]):
            chart_type = "bar"

        # Create sample data based on content context
        if "growth" in content_lower or "increase" in content_lower:
            return {
                "type": chart_type,
                "title": "Growth Metrics",
                "categories": ["Q1", "Q2", "Q3", "Q4"],
                "series": [
                    {"name": "Actual", "values": [15, 25, 40, 55]},
                    {"name": "Target", "values": [20, 30, 45, 60]},
                ],
            }
        if "market" in content_lower or "share" in content_lower:
            return {
                "type": "pie",
                "title": "Market Distribution",
                "categories": ["Company A", "Company B", "Company C", "Others"],
                "series": [{"name": "Market Share", "values": [35, 25, 20, 20]}],
            }
        if "performance" in content_lower or "metrics" in content_lower:
            return {
                "type": chart_type,
                "title": "Performance Metrics",
                "categories": ["Jan", "Feb", "Mar", "Apr", "May", "Jun"],
                "series": [{"name": "Performance", "values": [78, 82, 85, 88, 91, 95]}],
            }
        # Generic data
        return {
            "type": chart_type,
            "title": "Data Visualization",
            "categories": ["Category 1", "Category 2", "Category 3", "Category 4"],
            "series": [
                {"name": "Series 1", "values": [10, 20, 30, 40]},
                {"name": "Series 2", "values": [15, 25, 35, 45]},
            ],
        }

    def _set_text_preserving_formatting_with_layout(
        self, placeholder, text: str, layout_formatting: Dict[str, Any]
    ) -> None:
        """
        Set text while preserving ALL formatting from layout template including colors
        
        Args:
            placeholder: PowerPoint placeholder object
            text: Text content to set
            layout_formatting: Formatting from layout template
        """
        if not hasattr(placeholder, 'text_frame'):
            return
        
        text_frame = placeholder.text_frame
        
        # First, extract formatting from the layout using our extractor
        try:
            slide = placeholder.part.slide
            layout = slide.slide_layout
            placeholder_idx = placeholder.placeholder_format.idx
            
            # Extract layout formatting if not already done
            if not self.layout_formatter.placeholder_formats:
                self.layout_formatter.extract_from_layout(layout)
            
            # Get formatting for this specific placeholder
            extracted_formatting = self.layout_formatter.get_placeholder_formatting(placeholder_idx)
            if extracted_formatting:
                layout_formatting.update(extracted_formatting)
                print(f"  ✓ Extracted formatting for placeholder {placeholder_idx}: {extracted_formatting.get('name', 'unnamed')}")
            
            # Also try to get formatting directly from the layout placeholder
            for layout_ph in layout.placeholders:
                if layout_ph.placeholder_format.idx == placeholder_idx:
                    # Get font info directly from layout
                    if hasattr(layout_ph, 'text_frame') and layout_ph.text_frame:
                        if layout_ph.text_frame.paragraphs:
                            first_para = layout_ph.text_frame.paragraphs[0]
                            
                            # Capture bullet formatting
                            if hasattr(first_para, 'bullet'):
                                layout_formatting['bullet_char'] = first_para.bullet.char
                                layout_formatting['bullet_type'] = first_para.bullet.type
                                
                            # Get font from paragraph or runs
                            if first_para.runs:
                                run = first_para.runs[0]
                                if hasattr(run, 'font') and run.font:
                                    if run.font.name:
                                        layout_formatting['font_name'] = run.font.name
                                    if run.font.size:
                                        layout_formatting['font_size'] = run.font.size
                                    if run.font.bold is not None:
                                        layout_formatting['font_bold'] = run.font.bold
                                    if run.font.italic is not None:
                                        layout_formatting['font_italic'] = run.font.italic
                                    # IMPORTANT: Capture font color
                                    if hasattr(run.font, 'color') and run.font.color:
                                        try:
                                            if hasattr(run.font.color, 'rgb') and run.font.color.rgb:
                                                layout_formatting['font_color_rgb'] = run.font.color.rgb
                                            elif hasattr(run.font.color, 'theme_color'):
                                                layout_formatting['font_color_theme'] = run.font.color.theme_color
                                                layout_formatting['font_color_brightness'] = run.font.color.brightness
                                        except:
                                            pass
                            elif hasattr(first_para, 'font') and first_para.font:
                                if first_para.font.name:
                                    layout_formatting['font_name'] = first_para.font.name
                                if first_para.font.size:
                                    layout_formatting['font_size'] = first_para.font.size
                                if first_para.font.bold is not None:
                                    layout_formatting['font_bold'] = first_para.font.bold
                                if first_para.font.italic is not None:
                                    layout_formatting['font_italic'] = first_para.font.italic
                                # Capture font color from paragraph
                                if hasattr(first_para.font, 'color') and first_para.font.color:
                                    try:
                                        if hasattr(first_para.font.color, 'rgb') and first_para.font.color.rgb:
                                            layout_formatting['font_color_rgb'] = first_para.font.color.rgb
                                        elif hasattr(first_para.font.color, 'theme_color'):
                                            layout_formatting['font_color_theme'] = first_para.font.color.theme_color
                                            layout_formatting['font_color_brightness'] = first_para.font.color.brightness
                                    except:
                                        pass
                    break
        except Exception as e:
            print(f"  → Could not extract layout formatting: {e}")
        
        # Get placeholder index for formatting application
        try:
            placeholder_idx = placeholder.placeholder_format.idx
        except:
            placeholder_idx = 0
        
        # Apply markdown formatting while preserving template styling
        # Pass the layout formatting to the markdown formatter
        self.markdown_formatter.format_text_frame(text_frame, text, layout_formatting)
        
        # Now apply additional paragraph-level formatting from layout
        # This ensures alignment, bullets, and spacing are preserved
        if layout_formatting:
            print("  ✓ Applying complete formatting from layout")
            
            for paragraph in text_frame.paragraphs:
                # Apply paragraph-level formatting from the layout extractor
                self.layout_formatter.apply_paragraph_formatting(
                    paragraph, placeholder_idx
                )
                
                # Apply bullet formatting if needed
                if paragraph.level > 0 or layout_formatting.get('bullet'):
                    self.layout_formatter.apply_bullet_formatting(
                        paragraph, placeholder_idx
                    )
                
                # Apply text run formatting for each run in the paragraph
                for run in paragraph.runs:
                    self.layout_formatter.apply_formatting_to_run(run, placeholder_idx)
                
                # Apply to paragraph font if possible
                if hasattr(paragraph, 'font') and paragraph.font:
                    try:
                        if layout_formatting.get('font_name'):
                            paragraph.font.name = layout_formatting['font_name']
                    except:
                        pass
                    
                    # Apply bold/italic from template if not overridden by markdown
                    if layout_formatting.get('font_bold') is not None and paragraph.font.bold is None:
                        try:
                            paragraph.font.bold = layout_formatting['font_bold']
                        except:
                            pass
                    
                    if layout_formatting.get('font_italic') is not None and paragraph.font.italic is None:
                        try:
                            paragraph.font.italic = layout_formatting['font_italic']
                        except:
                            pass
                    
                    # Apply font color
                    if layout_formatting.get('font_color_rgb'):
                        try:
                            paragraph.font.color.rgb = layout_formatting['font_color_rgb']
                        except:
                            pass
                    # Skip theme colors - they can cause XML corruption if not handled properly
                    # Theme colors need special handling that python-pptx doesn't always support
                    pass
                
                # Apply to all runs
                for run in paragraph.runs:
                    if hasattr(run, 'font') and run.font:
                        # Always apply font name
                        try:
                            if layout_formatting.get('font_name'):
                                run.font.name = layout_formatting['font_name']
                        except:
                            pass
                        
                        # Apply size if available
                        if layout_formatting.get('font_size'):
                            try:
                                run.font.size = layout_formatting['font_size']
                            except:
                                pass
                        
                        # Apply base bold/italic if the run doesn't have markdown overrides
                        if run.font.bold is None and layout_formatting.get('font_bold') is not None:
                            try:
                                run.font.bold = layout_formatting['font_bold']
                            except:
                                pass
                        
                        if run.font.italic is None and layout_formatting.get('font_italic') is not None:
                            try:
                                run.font.italic = layout_formatting['font_italic']
                            except:
                                pass
                        
                        # IMPORTANT: Apply font color to runs
                        if layout_formatting.get('font_color_rgb'):
                            try:
                                from pptx.dml.color import RGBColor
                                run.font.color.rgb = layout_formatting['font_color_rgb']
                                print(f"    → Applied RGB color to run")
                            except Exception as e:
                                print(f"    → Could not apply RGB color: {e}")
                        # Skip theme colors - they can cause XML corruption
                        # python-pptx doesn't always handle theme colors correctly
                        pass

    def _set_text_preserving_formatting(self, text_frame, content: str) -> None:
        """
        Set text content with markdown formatting while preserving template styling

        Args:
            text_frame: PowerPoint text frame object
            content: Markdown-formatted text content to set
        """
        try:
            # First capture the CURRENT formatting from the placeholder
            # This is important because it may have template-specific fonts like Helvetica
            original_format = self._capture_original_formatting(text_frame)
            
            # Extract layout formatting if available
            layout_formatting = {}
            try:
                # Get the slide and layout to extract formatting
                for shape in text_frame._parent.part.slide.shapes:
                    if hasattr(shape, 'text_frame') and shape.text_frame == text_frame:
                        if hasattr(shape, 'placeholder_format'):
                            placeholder_idx = shape.placeholder_format.idx
                            layout = text_frame._parent.part.slide.slide_layout
                            
                            # Extract layout formatting if not already done
                            if not self.layout_formatter.placeholder_formats:
                                self.layout_formatter.extract_from_layout(layout)
                            
                            # Get formatting for this specific placeholder
                            layout_formatting = self.layout_formatter.get_placeholder_formatting(placeholder_idx)
                            break
            except:
                pass  # If we can't get layout formatting, continue without it
            
            # Use markdown formatter with layout formatting
            self.markdown_formatter.format_text_frame(text_frame, content, layout_formatting)
            
            # After markdown formatting, ensure we restore the original font name
            # if it was captured (markdown formatter might miss it)
            if original_format.get("font_name") and text_frame.paragraphs:
                for paragraph in text_frame.paragraphs:
                    for run in paragraph.runs:
                        try:
                            run.font.name = original_format["font_name"]
                        except:
                            pass  # Some fonts might be theme-controlled
            
            print("  ✓ Applied markdown formatting with font preservation")

        except Exception as e:
            print(f"Error applying markdown formatting: {e}")
            # Fallback to simple text setting
            text_frame.text = content

    def _capture_original_formatting(self, text_frame) -> dict:
        """
        Capture original formatting properties from the text frame

        Args:
            text_frame: PowerPoint text frame object

        Returns:
            Dictionary containing original formatting properties
        """
        original_format: Dict[str, Any] = {
            "font_name": None,
            "font_size": None,
            "bold": None,
            "italic": None,
            "color": None,
            "paragraph_format": None,
        }

        try:
            if text_frame.paragraphs:
                # Get formatting from first paragraph
                paragraph = text_frame.paragraphs[0]
                original_format["paragraph_format"] = {
                    "level": paragraph.level,
                    "alignment": (
                        paragraph.alignment if hasattr(paragraph, "alignment") else None
                    ),
                }

                # Get formatting from paragraph font
                if hasattr(paragraph, "font"):
                    para_font = paragraph.font
                    original_format["font_name"] = para_font.name
                    original_format["font_size"] = para_font.size
                    original_format["bold"] = para_font.bold
                    original_format["italic"] = para_font.italic
                    original_format["color"] = para_font.color

                # If paragraph font is None, try first run
                if paragraph.runs:
                    run = paragraph.runs[0]
                    if hasattr(run, "font"):
                        run_font = run.font
                        # Only override None values from paragraph
                        if original_format["font_name"] is None:
                            original_format["font_name"] = run_font.name
                        if original_format["font_size"] is None:
                            original_format["font_size"] = run_font.size
                        if original_format["bold"] is None:
                            original_format["bold"] = run_font.bold
                        if original_format["italic"] is None:
                            original_format["italic"] = run_font.italic
                        if original_format["color"] is None:
                            original_format["color"] = run_font.color

        except Exception:
            # If we can't capture formatting, return defaults
            pass

        return original_format

    def _add_formatted_content(
        self, text_frame, content: str, original_format: dict
    ) -> None:
        """
        Add content to text frame while applying original formatting

        Args:
            text_frame: PowerPoint text frame object
            content: Text content to add
            original_format: Original formatting properties to preserve
        """
        try:
            # Add a paragraph
            paragraph = (
                text_frame.paragraphs[0]
                if text_frame.paragraphs
                else text_frame.add_paragraph()
            )

            # Apply original paragraph formatting
            if original_format.get("paragraph_format"):
                para_format = original_format["paragraph_format"]
                if para_format.get("level") is not None:
                    paragraph.level = para_format["level"]
                if para_format.get("alignment") is not None:
                    paragraph.alignment = para_format["alignment"]

            # Set the text content
            paragraph.text = content

            # Apply original font formatting to runs if we have them
            if paragraph.runs and any(
                v is not None
                for v in [
                    original_format.get("font_name"),
                    original_format.get("font_size"),
                    original_format.get("bold"),
                    original_format.get("italic"),
                ]
            ):
                for run in paragraph.runs:
                    if hasattr(run, "font"):
                        font = run.font

                        # Apply captured formatting only if it was explicitly set
                        if original_format.get("font_name") is not None:
                            font.name = original_format["font_name"]
                        if original_format.get("font_size") is not None:
                            font.size = original_format["font_size"]
                        if original_format.get("bold") is not None:
                            font.bold = original_format["bold"]
                        if original_format.get("italic") is not None:
                            font.italic = original_format["italic"]
                        # Note: Color formatting is more complex and often theme-based
                        # so we'll let PowerPoint handle it via the template

        except Exception as e:
            print(f"Error applying formatting: {e}")
            # Fallback: just set the text
            if text_frame.paragraphs:
                text_frame.paragraphs[0].text = content
            else:
                text_frame.add_paragraph().text = content

    def _format_text_frame(self, text_frame) -> None:
        """
        Apply basic formatting to text frame (DEPRECATED - now preserved from template)

        Args:
            text_frame: PowerPoint text frame object
        """
        # This method is now deprecated as we preserve original formatting
        # Keeping it for backward compatibility but it does nothing
        pass

    def _ensure_output_path(self, output_path: str) -> str:
        """
        Ensure the output path has the correct extension and directory exists
        Uses centralized generated_presentations folder for organization

        Args:
            output_path: Requested output path

        Returns:
            Full output path with proper extension in generated_presentations folder
        """
        # Use default output directory for generated presentations
        output_dir = os.getenv("OUTPUT_DIRECTORY", "generated_presentations")

        # If output_path is relative and doesn't include the output directory,
        # prepend it
        is_relative = not os.path.isabs(output_path)
        missing_output_dir = not output_path.startswith(output_dir)
        if is_relative and missing_output_dir:
            output_path = os.path.join(output_dir, output_path)

        # Ensure .pptx extension
        if not output_path.lower().endswith(".pptx"):
            output_path += ".pptx"

        # Ensure directory exists
        output_dir_path = os.path.dirname(output_path)
        if output_dir_path and not os.path.exists(output_dir_path):
            os.makedirs(output_dir_path, exist_ok=True)
            print(f"📁 Created output directory: {output_dir_path}")

        return output_path

    def preview_content(self, topic: str) -> str:
        """
        Generate and return a text preview of the presentation content

        Args:
            topic: The presentation topic

        Returns:
            String preview of the presentation
        """
        return self.content_generator.create_content_preview(topic)

    def get_available_layouts(self) -> Dict[int, Dict[str, Any]]:
        """
        Get information about available layouts in the template

        Returns:
            Dictionary of layout information
        """
        return self.content_generator.analyze_template()

    def print_template_analysis(self) -> None:
        """Print analysis of the template layouts"""
        self.content_generator.layout_analyzer.print_layout_summary()
