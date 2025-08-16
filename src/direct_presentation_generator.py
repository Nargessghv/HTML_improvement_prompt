"""
Direct Presentation Generator

This module generates a complete presentation by adding slides directly to a single
presentation, avoiding the formatting issues that occur when copying between presentations.
"""

import os
import logging
from pathlib import Path
from typing import Dict, Any, List, Optional
from datetime import datetime

from pptx import Presentation

from .database import get_supabase_client
from .llm_client import SlideContent

logger = logging.getLogger(__name__)


class DirectPresentationGenerator:
    """
    Generates complete presentations by adding slides directly to a single presentation,
    preserving all template formatting and avoiding copying issues.
    """

    def __init__(self, template_path: str):
        """Initialize with template path"""
        self.template_path = template_path
        self.db = get_supabase_client()

    async def generate_complete_presentation(
        self,
        project_id: str,
        output_path: str,
        slide_contents: List[SlideContent],
        layouts_info: Optional[Dict[str, Any]] = None,
        dynamic_models: Optional[Dict[str, Any]] = None
    ) -> Dict[str, Any]:
        """
        Generate a complete presentation by adding all slides directly to one presentation
        
        This approach preserves ALL formatting because:
        1. All slides share the same template
        2. No copying between presentations
        3. Theme relationships are maintained
        4. Layout inheritance works properly
        
        Args:
            project_id: Project identifier
            output_path: Where to save the final presentation
            slide_contents: List of slide content objects
            layouts_info: Layout analysis information
            dynamic_models: Dynamic content models
            
        Returns:
            Result dictionary with success status and file info
        """
        try:
            logger.info(f"Generating complete presentation with {len(slide_contents)} slides")
            
            # Create presentation from template
            presentation = Presentation(self.template_path)
            
            # Clear any existing slides from template
            self._clear_template_slides(presentation)
            
            # Add all slides directly to this presentation
            slides_added = 0
            for i, slide_content in enumerate(slide_contents):
                try:
                    success = self._add_slide_directly(
                        presentation=presentation,
                        slide_content=slide_content,
                        slide_number=i + 1,
                        layouts_info=layouts_info,
                        dynamic_models=dynamic_models
                    )
                    
                    if success:
                        slides_added += 1
                        logger.info(f"✅ Added slide {i + 1} directly to presentation")
                    else:
                        logger.warning(f"⚠️ Failed to add slide {i + 1}")
                        
                except Exception as e:
                    logger.error(f"❌ Error adding slide {i + 1}: {e}")
                    continue
            
            if slides_added == 0:
                return {
                    "success": False,
                    "error": "No slides were successfully added",
                    "slides_added": 0
                }
            
            # Save the complete presentation
            presentation.save(output_path)
            
            # Get file information
            file_size = os.path.getsize(output_path)
            
            # Update database
            await self._update_project_final_presentation(
                project_id=project_id,
                file_path=output_path,
                file_size=file_size,
                slides_count=slides_added
            )
            
            logger.info(f"✅ Generated complete presentation: {slides_added}/{len(slide_contents)} slides")
            
            return {
                "success": True,
                "output_path": output_path,
                "slides_added": slides_added,
                "total_slides": len(slide_contents),
                "file_size": file_size,
                "message": f"Generated presentation with {slides_added} slides using direct approach"
            }
            
        except Exception as e:
            logger.error(f"❌ Failed to generate complete presentation: {e}")
            return {
                "success": False,
                "error": str(e),
                "slides_added": 0
            }

    def _clear_template_slides(self, presentation: Presentation):
        """Remove any existing slides from the template"""
        try:
            while len(presentation.slides) > 0:
                rId = presentation.slides._sldIdLst[0].rId
                presentation.part.drop_rel(rId)
                del presentation.slides._sldIdLst[0]
            logger.debug("Cleared template slides")
        except Exception as e:
            logger.warning(f"Could not clear template slides: {e}")

    def _add_slide_directly(
        self,
        presentation: Presentation,
        slide_content: SlideContent,
        slide_number: int,
        layouts_info: Optional[Dict[str, Any]] = None,
        dynamic_models: Optional[Dict[str, Any]] = None
    ) -> bool:
        """
        Add a slide directly to the presentation using the proven SlideGenerator approach
        
        This preserves formatting because the slide is created within the same presentation
        that contains the template and theme information.
        """
        try:
            # Get layout index from slide content
            layout_index = self._get_layout_index(slide_content)
            
            # Validate layout index
            if layout_index < 0 or layout_index >= len(presentation.slide_layouts):
                logger.warning(f"Invalid layout index {layout_index}, using default")
                layout_index = 1  # Use content layout
            
            # Create slide with proper layout
            slide_layout = presentation.slide_layouts[layout_index]
            slide = presentation.slides.add_slide(slide_layout)
            
            # Apply content using the proven SlideGenerator approach
            self._apply_content_to_slide(
                slide=slide,
                slide_content=slide_content,
                layouts_info=layouts_info,
                dynamic_models=dynamic_models,
                slide_number=slide_number
            )
            
            return True
            
        except Exception as e:
            logger.error(f"Error adding slide {slide_number}: {e}")
            return False

    def _apply_content_to_slide(
        self,
        slide,
        slide_content: SlideContent,
        layouts_info: Optional[Dict[str, Any]] = None,
        dynamic_models: Optional[Dict[str, Any]] = None,
        slide_number: int = 1
    ):
        """
        Apply content to slide using the proven SlideGenerator methods
        
        This maintains formatting because we're working within the same presentation context.
        """
        try:
            # Import here to avoid circular imports
            from .slide_generator import SlideGenerator
            
            # Create temporary SlideGenerator to use its proven methods
            slide_generator = SlideGenerator(self.template_path)
            
            # Process content
            content = slide_content.content if hasattr(slide_content, 'content') else {}
            
            # Auto-add LOCKED_ backgrounds if available
            content = self._process_locked_placeholders(content, slide)
            
            # Use the proven content application methods
            if layouts_info and hasattr(slide_generator, 'content_generator'):
                # Set up layouts_info for proper placeholder mapping
                if not hasattr(slide_generator.content_generator, 'layouts_info'):
                    class SimpleContentGenerator:
                        def __init__(self, layouts_info):
                            self.layouts_info = layouts_info
                    slide_generator.content_generator = SimpleContentGenerator(layouts_info)
                else:
                    slide_generator.content_generator.layouts_info = layouts_info
            
            # Apply content using the proven placeholder population method
            slide_generator._populate_slide_placeholders(slide, content)
            
            logger.debug(f"Applied content to slide {slide_number}")
            
        except Exception as e:
            logger.error(f"Error applying content to slide: {e}")

    def _process_locked_placeholders(self, content: Dict[str, str], slide) -> Dict[str, str]:
        """Process LOCKED_ placeholders by adding PNG paths"""
        try:
            template_folder = Path(self.template_path).parent
            
            # Check slide layout for LOCKED_ placeholders
            slide_layout = slide.slide_layout
            for layout_ph in slide_layout.placeholders:
                ph_name = layout_ph.name
                if ph_name and ph_name.startswith("LOCKED_"):
                    png_path = template_folder / f"{ph_name}.png"
                    if png_path.exists() and ph_name not in content:
                        content[ph_name] = str(png_path)
                        logger.debug(f"Auto-added LOCKED_ background: {ph_name}")
            
            return content
            
        except Exception as e:
            logger.warning(f"Error processing LOCKED_ placeholders: {e}")
            return content

    def _get_layout_index(self, slide_content: SlideContent) -> int:
        """Extract layout index from slide content"""
        try:
            if hasattr(slide_content, 'layout_index'):
                return slide_content.layout_index
            
            if hasattr(slide_content, 'content') and isinstance(slide_content.content, dict):
                if 'layout_index' in slide_content.content:
                    return int(slide_content.content['layout_index'])
            
            # Default to content layout
            return 1
            
        except Exception:
            return 1

    async def _update_project_final_presentation(
        self,
        project_id: str,
        file_path: str,
        file_size: int,
        slides_count: int
    ):
        """Update database with final presentation information"""
        try:
            # Update projects table
            update_data = {
                "final_presentation_path": file_path,
                "final_presentation_size": file_size,
                "final_presentation_slides": slides_count,
                "final_presentation_generated_at": datetime.now().isoformat(),
                "updated_at": datetime.now().isoformat()
            }
            
            self.db.client.table("projects").update(update_data).eq("id", project_id).execute()
            
            logger.info(f"Updated project {project_id} with final presentation info")
            
        except Exception as e:
            logger.error(f"Failed to update project with final presentation info: {e}")


class DirectSlideGenerator:
    """
    Alternative to IndividualSlideGenerator that creates slides directly in the final presentation
    """
    
    def __init__(self, template_path: str):
        self.template_path = template_path
        self.direct_generator = DirectPresentationGenerator(template_path)
    
    async def generate_presentation_directly(
        self,
        project_id: str,
        output_path: str,
        slide_contents: List[SlideContent],
        layouts_info: Optional[Dict[str, Any]] = None,
        dynamic_models: Optional[Dict[str, Any]] = None
    ) -> Dict[str, Any]:
        """
        Generate complete presentation directly without individual files
        
        This is the recommended approach for avoiding formatting issues.
        """
        return await self.direct_generator.generate_complete_presentation(
            project_id=project_id,
            output_path=output_path,
            slide_contents=slide_contents,
            layouts_info=layouts_info,
            dynamic_models=dynamic_models
        )