from pptx import Presentation


def analyze_placeholder_formatting(template_path):
    """Analyze detailed formatting of placeholders in template"""
    prs = Presentation(template_path)

    print("=== TEMPLATE FORMATTING ANALYSIS ===\n")

    for layout_idx, layout in enumerate(prs.slide_layouts):
        print(f"Layout {layout_idx}: {layout.name}")
        print("-" * 50)

        for placeholder in layout.placeholders:
            print(
                f"Placeholder: '{placeholder.name}' (Index: {placeholder.placeholder_format.idx})"
            )
            print(f"Type: {placeholder.placeholder_format.type}")

            # Check if it has text frame
            if hasattr(placeholder, "text_frame") and placeholder.text_frame:
                text_frame = placeholder.text_frame
                print("Has text frame: Yes")

                # Check paragraphs
                for para_idx, paragraph in enumerate(text_frame.paragraphs):
                    print(f"  Paragraph {para_idx}:")
                    print(f"    Text: '{paragraph.text}'")
                    print(f"    Level: {paragraph.level}")

                    # Check paragraph font
                    if paragraph.font:
                        font = paragraph.font
                        print("    Paragraph Font:")
                        print(f"      Name: {font.name}")
                        print(f"      Size: {font.size}")
                        print(f"      Bold: {font.bold}")
                        print(f"      Italic: {font.italic}")

                    # Check runs
                    for run_idx, run in enumerate(paragraph.runs):
                        print(f"    Run {run_idx}:")
                        print(f"      Text: '{run.text}'")
                        if run.font:
                            font = run.font
                            print(f"      Font Name: {font.name}")
                            print(f"      Font Size: {font.size}")
                            print(f"      Bold: {font.bold}")
                            print(f"      Italic: {font.italic}")
                            print(f"      Color: {font.color}")
            else:
                print("Has text frame: No")

            print()

        print("=" * 50)
        print()


if __name__ == "__main__":
    analyze_placeholder_formatting("template.pptx")
